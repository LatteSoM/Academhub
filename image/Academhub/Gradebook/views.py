import logging
import re
from collections import defaultdict

from django.contrib import messages
from django.contrib.auth.decorators import permission_required
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse, FileResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views import View
from openpyxl.styles import PatternFill, Font, Alignment, Side, Border
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from unidecode import unidecode

from Academhub.generic import BulkUpdateView, ObjectTableView, ObjectDetailView, ObjectUpdateView, ObjectCreateView, \
    ObjectTemplateView
from Academhub.models import GradebookStudents, Gradebook, CustomUser, SubTable, Button, Student, \
    GroupStudents
from Academhub.modules.documentGenPars import GradebookDocumentGenerator
from Academhub.utils import getpermission, getpattern
from Gradebook.filters import *
from Gradebook.filters import GradeBookTeachersFilter
from Gradebook.forms import *
from Gradebook.forms import GenerateGradebookForm, GetStatisticksGradebookForm, GetDiplomTypeGradebookForm
from Gradebook.mixins import GradeBookMixin
from Gradebook.tables import *
from Gradebook.tables import TeacherGradeBookTable

#
# Create your views here.
#

__all__ = (
    'download_report',
    'export_grades',
    "export_diplom_types",
    'GradebookTableView',
    'GradebookClosedList',
    'GradebookDetailView',
    'GradebookCreateView',
    'GradebookUpdateView',
    'GradebookGenerateView',
    'DiplomTypeIdentifier',
    'GradebookRedirectView',
    'ViewStudentsGradesForSemester',
    'check_and_open_gradebook',
    'TeachersGradeBookTableView',
    'GradebookStudentBulkUpdateView',
)

from Gradebook.utils import (transliterate, get_disciplines_with_last_semester, check_students_for_academic_debts,
                             perform_generation,
                             fetch_group_semester_attestation_table, )


class GradebookStudentBulkUpdateView(BulkUpdateView):
    model = GradebookStudents
    form_class = GradebookStudentsForm
    template_name = 'Gradebook/create/grade_book_students.html'

    def dispatch(self, request, *args, **kwargs):
        self.gradebook_pk = self.kwargs.get('pk', None)
        return super().dispatch(request, *args, **kwargs)

    def get_queryset(self):
        return self.model.objects.filter(gradebook__pk = self.gradebook_pk)

    def get(self, request, *args, **kwargs):
        self.gradebook = get_object_or_404(Gradebook, pk=self.gradebook_pk)
        if self.gradebook.status not in ["Открыта", "Заполнена"]:
            messages.error(
                request,
                "Заполнить/изменить заполнение можно только открытую или уже заполненную ведомость"
            )
            logger.warning(
                f"Attempt to access gradebook {self.gradebook_pk} with status {self.gradebook.status}"
            )
            return redirect(reverse('gradebook_detail', kwargs={'pk': self.gradebook_pk}))
        return super().get(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['gradebook'] = self.gradebook
        return context

    def save_form(self, request):
        form = super().save_form(request)

        if form.is_valid():
            gradebook = get_object_or_404(
                Gradebook,
                pk=self.gradebook_pk
            )

            check_students_for_academic_debts(gradebook.group, gradebook.semester_number)

            gradebook.status = Gradebook.STATUS_CHOICE[2][1]
            gradebook.date_of_filling = timezone.now()
            gradebook.save()

        return form

    def post(self, request, *args, **kwargs):
        formset = self.save_form(request)

        if formset.is_valid():
            user = get_object_or_404(CustomUser, pk=self.request.user.pk)

            if user.is_teacher:
                response = redirect('gradebookteacher_list')
            else:
                response = redirect('gradebook_list')
            response['Location'] += '?is_filled=true'
            return response
        return super().post(request, *args, **kwargs)


class GradebookGenerateView(ObjectTemplateView):
    """Отдельный класс для обработки генерации"""
    template_name = 'Gradebook/inc/generate_form.html'

    permission_required = getpermission('Gradebook', 'create_gradebook')

    def get(self, request):
        form = GenerateGradebookForm()
        return render(request, self.template_name, {'form': form})

    def post(self, request):
        form = GenerateGradebookForm(request.POST)  # Создаем экземпляр формы

        if form.is_valid():  # Проверяем валидность ФОРМЫ, а не представления
            try:
                semester = form.cleaned_data['semester']  # Данные берем из формы
                gradebook_type = form.cleaned_data['type_of_gradebook']
                result = perform_generation(semester, gradebook_type)

                # Формируем базовое сообщение
                success_part = []
                if result['generated'] > 0:
                    success_part.append(f"сгенерировано {result['generated']}")
                if result['updated'] > 0:
                    success_part.append(f"обновлено {result['updated']}")

                # Проверяем, есть ли успешные действия
                if success_part:
                    success_message = " и ".join(success_part)
                    message = f"Успешно {success_message} ведомостей для {semester} семестра! С типом {result['grade_book_type']}. Ошибок: {result['errors']}"

                    # Выбираем тип сообщения в зависимости от наличия ошибок
                    if result['errors'] > 0:
                        messages.warning(request, message)
                    else:
                        messages.success(request, message)
                else:
                    messages.error(request, "Не удалось сгенерировать ни одной ведомости!")

                return redirect('gradebook_list')

            except Exception as e:
                messages.error(request, f"Ошибка генерации: {str(e)}")
                return redirect('gradebook_list')

        # Если форма невалидна, показываем ее с ошибками
        return render(request, self.template_name, {'form': form})


class ViewStudentsGradesForSemester(ObjectTemplateView):
    """Отдельный класс для обработки статистической таблицы студентов"""
    template_name = 'Gradebook/inc/statisticks_form.html'
    permission_required = getpermission('Gradebook', 'view_student_statick')

    def get(self, request):
        form = GetStatisticksGradebookForm()
        return render(request, self.template_name, {'form': form})

    def post(self, request):
        form = GetStatisticksGradebookForm(request.POST)  # Создаем экземпляр формы

        if form.is_valid():  # Проверяем валидность формы
            try:
                semester = form.cleaned_data['semester']  # Данные берем из формы
                group = form.cleaned_data['group']

                # Получаем данные с помощью функции
                data = fetch_group_semester_attestation_table(semester, group)

                # Проверяем, найден ли учебный план
                if data is None:
                    messages.error(request, "Учебный план для данной группы не найден.")
                    return render(request, self.template_name, {'form': form})

                # Передаем все данные в шаблон
                return render(request, 'Gradebook/list/group_semester_grades.html', {
                    'students': data['students'],
                    'disciplines': data['disciplines'],
                    'grades': data['grades'],
                    'semestr': data['semestr'],
                    'group': data['group'],
                    'curriculum_items': data['curriculum_items'],
                    'quality_percentages': data['quality_percentages'],
                    'success_percentages': data['success_percentages'],
                    'average_grades': data['average_grades'],
                    'total_quality_percentage': data['total_quality_percentage'],
                    'total_success_percentage': data['total_success_percentage'],
                    'total_average_grade': data['total_average_grade'],
                })

            except Exception as e:
                print(e)
                messages.error(request, "Не удалось сгенерировать статистику!")
                return render(request, self.template_name, {'form': form})

        # Если форма невалидна, показываем ее с ошибками
        return render(request, self.template_name, {'form': form})


def export_grades(request):
    try:
        # Получение и валидация параметров
        semester = request.GET.get('semester')
        group_id = request.GET.get('group')

        if not semester or not group_id:
            raise ValueError("Не указаны обязательные параметры")

        group = get_object_or_404(GroupStudents, id=int(group_id))
        semester_number = int(semester)

    except (ValueError, TypeError) as e:
        return HttpResponse(f"Ошибка параметров: {str(e)}", status=400)
    except Exception as e:
        return HttpResponse(f"Ошибка сервера: {str(e)}", status=500)

    # Получение данных
    data = fetch_group_semester_attestation_table(semester_number, group)

    # Создание Excel-документа
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводная ведомость"

    # Стили оформления
    thin_side = Side(style='thin')
    thin_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side
    )
    center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    header_fill = PatternFill(start_color='D9E1F2', fill_type='solid')
    header_font = Font(bold=True)
    footer_fill = PatternFill(start_color='F0F0F0', fill_type='solid')

    # Цвета заливки
    FILLS = {
        'Первичная сдача': PatternFill(start_color='D4EDDA', fill_type='solid'),
        'Пересдача': PatternFill(start_color='CCE5FF', fill_type='solid'),
        'Комиссия': PatternFill(start_color='E2D9F3', fill_type='solid'),
        'неуд': PatternFill(start_color='FFC7CE', fill_type='solid')
    }

    # Заголовок документа
    ws.merge_cells('A1:M1')
    title_cell = ws['A1']
    title_cell.value = f"Сводная ведомость успеваемости студентов\nГруппа {group.full_name}\nСеместр {semester}"
    title_cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    title_cell.font = Font(bold=True, size=14)
    ws.row_dimensions[1].height = 60

    ws.cell(row=5, column=1, value="№ п/п").font = header_font
    ws.cell(row=5, column=2, value="Ф.И.О. студента").font = header_font
    ws.cell(row=5, column=3, value="Основа обучения").font = header_font

    # Группировка дисциплин по типам аттестации из Curriculum
    TYPE_MAPPING = {
        'Экзамен': 'Экзамены',
        'Зачет с оценкой': 'Диф. зачёты',
        'Курсовая работа': 'Курсовые работы',
        'Другие формы контроля': 'Другие формы'
    }

    discipline_groups = defaultdict(list)
    for curriculum_item in data['curriculum_items']:
        if curriculum_item.term != semester_number:
            continue
        discipline = curriculum_item.discipline
        category = TYPE_MAPPING.get(
            curriculum_item.code_of_type_work,
            'Другие формы'
        )
        if discipline.id not in {d.id for d in discipline_groups[category]}:
            discipline_groups[category].append(discipline)

    # Формирование структуры таблицы
    headers = ['№ п/п', 'Ф.И.О. студента', 'Основа обучения']
    column_map = {}
    current_col = 4

    categories_order = [
        'Экзамены',
        'Диф. зачёты',
        'Курсовые работы',
        'Другие формы'
    ]

    for category in categories_order:
        disciplines = sorted(
            discipline_groups.get(category, []),
            key=lambda x: x.name
        )
        if not disciplines:
            continue
        end_col = current_col + len(disciplines) - 1
        ws.merge_cells(
            start_row=4,
            start_column=current_col,
            end_row=4,
            end_column=end_col
        )
        cell = ws.cell(row=4, column=current_col, value=category)
        cell.alignment = center_alignment
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        for idx, disc in enumerate(disciplines, start=current_col):
            cell = ws.cell(row=5, column=idx, value=disc.name)
            cell.alignment = center_alignment
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin_border
            column_map[disc.id] = idx
        current_col += len(disciplines)

    # Заполнение данных студентов
    total_rows = len(data['students'])
    for row_idx, student in enumerate(data['students'], start=6):
        ws.cell(row=row_idx, column=1, value=row_idx - 5).alignment = center_alignment
        ws.cell(row=row_idx, column=1).border = thin_border
        ws.cell(row=row_idx, column=2, value=student.full_name).border = thin_border
        ws.cell(row=row_idx, column=3, value=student.education_basis).alignment = center_alignment
        ws.cell(row=row_idx, column=3).border = thin_border
        for category in categories_order:
            for disc in discipline_groups.get(category, []):
                col = column_map.get(disc.id)
                if not col:
                    continue
                grade_entry = data['grades'].get(student.id, {}).get(disc.id)
                cell = ws.cell(row=row_idx, column=col)
                cell.alignment = center_alignment
                cell.border = thin_border
                if grade_entry:
                    if grade_entry == "-":
                        cell.value = grade_entry
                    else:
                        cell.value = grade_entry.grade
                        if str(grade_entry.grade).lower() in ['неудовлетворительно', 'неявка', 'н/я', '2']:
                            cell.fill = FILLS['неуд']
                        elif grade_entry.gradebook.type_of_grade_book in FILLS:
                            cell.fill = FILLS[grade_entry.gradebook.type_of_grade_book]

    # Добавление метрик по дисциплинам
    metrics_start_row = total_rows + 6
    quality_row = metrics_start_row
    success_row = metrics_start_row + 1
    average_row = metrics_start_row + 2

    # Процент качества
    ws.merge_cells(start_row=quality_row, start_column=1, end_row=quality_row, end_column=3)
    quality_cell = ws.cell(row=quality_row, column=1, value="Процент качества")
    quality_cell.alignment = center_alignment
    quality_cell.font = header_font
    quality_cell.fill = footer_fill
    quality_cell.border = thin_border
    max_col = max(column_map.values()) if column_map else 4
    for discipline in data['disciplines']:
        col = column_map.get(discipline.id)
        if col:
            cell = ws.cell(row=quality_row, column=col, value=f"{data['quality_percentages'].get(discipline.id, 0):.2f}%")
            cell.alignment = center_alignment
            cell.border = thin_border
            cell.fill = footer_fill

    # Процент успеваемости
    ws.merge_cells(start_row=success_row, start_column=1, end_row=success_row, end_column=3)
    success_cell = ws.cell(row=success_row, column=1, value="Процент успеваемости")
    success_cell.alignment = center_alignment
    success_cell.font = header_font
    success_cell.fill = footer_fill
    success_cell.border = thin_border
    for discipline in data['disciplines']:
        col = column_map.get(discipline.id)
        if col:
            cell = ws.cell(row=success_row, column=col, value=f"{data['success_percentages'].get(discipline.id, 0):.2f}%")
            cell.alignment = center_alignment
            cell.border = thin_border
            cell.fill = footer_fill

    # Средний балл
    ws.merge_cells(start_row=average_row, start_column=1, end_row=average_row, end_column=3)
    average_cell = ws.cell(row=average_row, column=1, value="Средний балл")
    average_cell.alignment = center_alignment
    average_cell.font = header_font
    average_cell.fill = footer_fill
    average_cell.border = thin_border
    for discipline in data['disciplines']:
        col = column_map.get(discipline.id)
        if col:
            cell = ws.cell(row=average_row, column=col, value=f"{data['average_grades'].get(discipline.id, 0):.2f}")
            cell.alignment = center_alignment
            cell.border = thin_border
            cell.fill = footer_fill

    # Добавление групповых метрик в самом конце
    group_metrics_start_row = average_row + 2
    if total_rows > 0:
        # Процент качества по учебной группе
        ws.merge_cells(start_row=group_metrics_start_row, start_column=1, end_row=group_metrics_start_row, end_column=1)
        group_quality_cell = ws.cell(row=group_metrics_start_row, column=1, value="Процент качества по учебной группе")
        group_quality_cell.alignment = center_alignment
        group_quality_cell.font = header_font
        group_quality_cell.fill = footer_fill
        group_quality_cell.border = thin_border
        group_quality_value_cell = ws.cell(row=group_metrics_start_row, column=2, value=f"{data.get('total_quality_percentage', 0):.2f}%")
        group_quality_value_cell.alignment = center_alignment
        group_quality_value_cell.font = header_font
        group_quality_value_cell.fill = footer_fill
        group_quality_value_cell.border = thin_border

        # Процент успеваемости по группе
        ws.merge_cells(start_row=group_metrics_start_row + 1, start_column=1, end_row=group_metrics_start_row + 1, end_column=1)
        group_success_cell = ws.cell(row=group_metrics_start_row + 1, column=1, value="Процент успеваемости по группе")
        group_success_cell.alignment = center_alignment
        group_success_cell.font = header_font
        group_success_cell.fill = footer_fill
        group_success_cell.border = thin_border
        group_success_value_cell = ws.cell(row=group_metrics_start_row + 1, column=2, value=f"{data.get('total_success_percentage', 0):.2f}%")
        group_success_value_cell.alignment = center_alignment
        group_success_value_cell.font = header_font
        group_success_value_cell.fill = footer_fill
        group_success_value_cell.border = thin_border

        # Средний балл по группе
        ws.merge_cells(start_row=group_metrics_start_row + 2, start_column=1, end_row=group_metrics_start_row + 2, end_column=1)
        group_average_cell = ws.cell(row=group_metrics_start_row + 2, column=1, value="Средний балл по группе")
        group_average_cell.alignment = center_alignment
        group_average_cell.font = header_font
        group_average_cell.fill = footer_fill
        group_average_cell.border = thin_border
        group_average_value_cell = ws.cell(row=group_metrics_start_row + 2, column=2, value=f"{data.get('total_average_grade', 0):.2f}")
        group_average_value_cell.alignment = center_alignment
        group_average_value_cell.font = header_font
        group_average_value_cell.fill = footer_fill
        group_average_value_cell.border = thin_border

    # Настройка ширины столбцов
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = 20 if col_idx > 3 else 15

    # Применение границ только к используемым ячейкам
    max_col = max(column_map.values()) if column_map else 4  # Максимальный столбец с данными
    for row in ws.iter_rows(min_row=4, max_row=group_metrics_start_row + 2, min_col=1, max_col=max_col):
        for cell in row:
            if cell.row < group_metrics_start_row - 1:
                cell.border = thin_border

            if cell.row > 5 and cell.column in [1, 3]:
                cell.alignment = center_alignment

    # Формирование имени файла
    group_name_clean = unidecode(group.full_name)
    group_name_clean = re.sub(r'\W+', '_', group_name_clean)[:50]
    filename = f"Vedomost_{group_name_clean}_sem{semester}.xlsx"

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


class DiplomTypeIdentifier(ObjectTemplateView):
    template_name = "Gradebook/inc/diplom_type_form.html"

    def get(self, request):
        form = GetDiplomTypeGradebookForm()
        return render(request, self.template_name, {'form': form})

    def post(self, request):
        form = GetDiplomTypeGradebookForm(request.POST)

        if form.is_valid():
            try:
                form_group = form.cleaned_data['group']
                group = GroupStudents.objects.get(id=form_group.id)
                logger.info(f"Processing group: {group.full_name}")

                # Получаем дисциплины и общую сумму часов
                disciplines_with_semesters, total_all_hours = get_disciplines_with_last_semester(group)
                if not disciplines_with_semesters:
                    logger.warning(f"No disciplines found for group {group.full_name}")
                    messages.error(request, "Для выбранной группы нет дисциплин")
                    return render(request, self.template_name, {'form': form})

                logger.info(f"Disciplines: {[(d[0].name, d[1], d[2], d[3]) for d in disciplines_with_semesters]}")
                semesters_order = sorted({d[1] for d in disciplines_with_semesters})

                context = {
                    'group': group,
                    'students': Student.objects.filter(group=group, is_expelled=False, is_in_academ=False),
                    'disciplines_with_semesters': disciplines_with_semesters,
                    'semesters_order': semesters_order,
                    'total_all_hours': total_all_hours,
                    'current_year': timezone.now().year,
                    'education_basis': 'После 11 класса' if group.education_base == 'Среднее общее' else 'После 9 класса',
                    'total_disciplines': len(disciplines_with_semesters)
                }

                return render(request, 'Gradebook/list/diplom_type_identifier.html', context)
            except Exception as e:
                logger.error(f"Error calculating diploma types: {str(e)}", exc_info=True)
                messages.error(request, f"Не удалось рассчитать тип дипломов: {str(e)}")
        else:
            logger.warning(f"Invalid form data: {form.errors}")
            messages.error(request, "Некорректные данные формы")

        return render(request, self.template_name, {'form': form})


class GradebookTableView(ObjectTableView):
    """
    Класс для отображения таблицы ведомостей.
    """
    table_class = GradebookTable
    filterset_class = GradebookFilter
    queryset = Gradebook.objects.exclude(status=Gradebook.STATUS_CHOICE[3][1])

    permission_required = getpermission('Gradebook', 'view_gradebook')

    buttons = [
        Button (
            id='add',
            name = 'Добавить',
            link_name = getpattern(Gradebook, 'add'),
            permission = getpermission('Gradebook', 'create_gradebook'),
        ),
        Button(
            id = 'generate',
            name = 'Сгенерировать/Обновить ведомости',
            link_name = 'gradebook_generate',
            permission = getpermission('Gradebook', 'create_gradebook'),
        )
    ]

    def get_table_class(self):
        if self.request.GET.get("mobile") == "1":
            return GradebookMobileTable
        return GradebookTable

    def get(self, request, *args, **kwargs):

        is_closed = request.GET.get('is_closed', None)
        is_filled = request.GET.get('is_filled', None)

        if is_closed:
            messages.success(request, "Ведомость успешно закрыта!")
        if is_filled:
            messages.success(request, "Ведомость успешно заполнена!")

        base_url = reverse('gradebook_list')
        query_params = request.GET.copy()
        if 'page' in query_params:
            del query_params['page']
        if query_params:
            request.session['gradebook_return_to'] = f"{base_url}?{query_params.urlencode()}"
        else:
            request.session['gradebook_return_to'] = base_url


        return super().get(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['form'] = GenerateGradebookForm()  # Доп. форма
        return context


class GradebookClosedList(ObjectTableView):
    table_class = GradebookTable
    filterset_class = ClosedGradebookFilter
    queryset = Gradebook.objects.filter(status=Gradebook.STATUS_CHOICE[3][1])
    properties = {}

    permission_required = getpermission('Gradebook', 'view_gradebook')

    def get_table_class(self):
        if self.request.GET.get("mobile") == "1":
            return GradebookMobileTable
        return GradebookTable

    def get(self, request, *args, **kwargs):
        base_url = reverse('gradebookclosed_list')
        query_params = request.GET.copy()
        if 'page' in query_params:
            del query_params['page']
        if query_params:
            request.session['gradebook_return_to'] = f"{base_url}?{query_params.urlencode()}"
        else:
            request.session['gradebook_return_to'] = base_url
        return super().get(request, *args, **kwargs)

class TeachersGradeBookTableView(ObjectTableView):
    table_class = TeacherGradeBookTable
    filterset_class = GradeBookTeachersFilter
    queryset = Gradebook.objects.all()

    def get_queryset(self):
        # Получаем залогиненного пользователя
        user = get_object_or_404(CustomUser, pk=self.request.user.pk)

        # Проверка аутентификации и прав преподавателя
        if not user.is_authenticated or not hasattr(user, 'is_teacher') or not user.is_teacher:
            return Gradebook.objects.none()

        queryset = user.gradebooks.filter(status=Gradebook.STATUS_CHOICE[1][1])

        return queryset
 
    def get_table_class(self):
        if self.request.GET.get("mobile") == "1":
            return GradebookMobileTable
        return TeacherGradeBookTable

    def get(self, request, *args, **kwargs):

        is_filled = request.GET.get('is_filled', None)

        if is_filled:
            messages.success(request, "Ведомость успешно заполнена!")

        base_url = reverse('gradebookteacher_list')
        query_params = request.GET.copy()
        if 'page' in query_params:
            del query_params['page']
        if query_params:
            request.session['gradebook_return_to'] = f"{base_url}?{query_params.urlencode()}"
        else:
            request.session['gradebook_return_to'] = base_url

        return super().get(request, *args, **kwargs)


class GradebookRedirectView(View):
    def get(self, request, *args, **kwargs):
        # Получаем URL из сессии, по умолчанию редиректим на 'gradebook:list'
        return_to = request.session.get('gradebook_return_to', reverse('gradebook_list'))
        return redirect(return_to)

class GradebookDetailView(ObjectDetailView):
    """
    Класс для отображения детальной информации об учебном журнале.
    """
    model= Gradebook
    paginate_by   = 30
    template_name = 'Gradebook/detail/gradebook_student.html'
    permission_required = getpermission('Gradebook', 'view_gradebook')

    fieldset = {
        'Основная информация': [
            'name', 'discipline','status', 'number', 'type_of_grade_book', 'group', 'semester_number'
        ],
        'Дополнительная информация': [
            'date_of_opening', 'date_of_filling', 'date_of_closing',
            'amount_of_days_for_closing'
        ]
    }

    def has_permission(self):
        if super().has_permission():
            return True

        try:
            obj = self.get_object()
            if self.request.user.gradebooks.filter(id=obj.id).exists():
                return True
            return False
        except (AttributeError, PermissionDenied):
            return False

    def dispatch(self, request, *args, **kwargs):
        self.request = request
        return super().dispatch(request, *args, **kwargs)


    # Формируем список кнопок динамически
    buttons = [
        Button (
            id = 'change',
            name = 'Изменить',
            link_params = ['pk'],
            link_name = getpattern(Gradebook, 'change'),
            permission = getpermission('Gradebook', 'update_gradebook'),
        ),
        Button (
            id = 'to_list',
            name = 'К таблице',
            link_name = 'redirect_to_list',
            permission = getpermission('Gradebook', 'view_gradebook')
        ),
        Button (
            link_params = ['pk'],
            name = 'Открыть ведомость',
            link_name = 'check_open_gradebook',
            condition = lambda obj, user: obj.status == "Не заполнена",
            permission = getpermission('Gradebook', 'open_gradebook'),
        ),
        Button (
            link_params = ['pk'],
            name = 'Закрыть ведомость',
            link_name = 'download_report',
            condition = lambda obj, user: obj.status == "Заполнена",
            permission = getpermission('Gradebook', 'close_gradebook'),
        ),
        Button (
            link_params = ['pk'],
            name = 'Скачать документ',
            link_name = 'download_report',
            condition = lambda obj, user: obj.status == "Закрыта",
            permission = getpermission('Gradebook', 'download_gradebook'),
        )
    ]

    def grade_book_student_filter(self, queryset):
        if queryset.filter(gradebook_id=self.pk).count() == 0:
            return GradebookStudents.objects.filter(gradebook=self.pk)
        return queryset.filter(gradebook__id=self.pk)

    def grade_book_teacher_filter(self, queryset):
        return self.teachers.all()

    tables = [
        SubTable (
            name='Студенты',
            table=GradebookStudentsTable,
            queryset=GradebookStudents.objects.all(),
            filter_func=grade_book_student_filter,
            buttons=[
                Button (
                    name = 'Заполнить ведомость',
                    link_name = 'gradebookstudents_bulk_create',
                    link_params = ['pk'],
                    condition= lambda obj, user: obj.status == "Открыта"
                )
            ]
        ),
        SubTable(
            name='Учителя',
            queryset=CustomUser.objects.all(),
            table=GradebookTeachersTable,
            filter_func=grade_book_teacher_filter
        )
    ]

    def get(self, request, *args, **kwargs):
        gradebook = self.get_object()

        is_active = request.GET.get('is_active',None)

        # Проверка основных полей
        required_fields = [
            gradebook.group_id,  # Проверка ForeignKey через id
            gradebook.discipline_id,  # Проверка ForeignKey дисциплины
            gradebook.semester_number,  # Проверка IntegerField
            gradebook.status,  # Проверка CharField статуса
        ]

        # Проверка текстовых полей на непустые значения
        text_fields_valid = all([
            gradebook.status and gradebook.status.strip(),
        ])

        relations_valid = all([
            gradebook.teachers.exists(),  # Проверка наличия преподавателей
            gradebook.students.exists(),  # Проверка наличия студентов
        ])

        # Комплексная проверка всех условий
        if not all(required_fields) or not text_fields_valid or not relations_valid and gradebook.status == "Не заполнена":
            messages.error(
                request,
                "Невозможно открыть ведомость! Заполните все обязательные поля: "
                "группа, дисциплина, название, статус, семестр, преподаватели и студенты."
            )
        else:
            if is_active:
                messages.success(request, "Ведомость успешно открыта!")

        return super().get(request, *args, **kwargs)

class GradebookUpdateView(GradeBookMixin, ObjectUpdateView):
    """
    Класс для обновления информации об учебном журнале.
    """
    form_class = GradebookForm
    queryset = Gradebook.objects.all()
    properties = {}

    permission_required = getpermission('Gradebook', 'update_gradebook')

    buttons = [
        Button (
            id='to_object',
            name = 'К объекту',
            link_params = ['pk'],
            link_name = getpattern(Gradebook, 'detail'),
            permission = getpermission('Gradebook', 'view_gradebook'),
        ),
        Button (
            id = 'to_list',
            name = 'К таблице',
            link_name = getpattern(Gradebook, 'list'),
            permission = getpermission('Gradebook', 'view_gradebook'),
        )
    ]


class GradebookCreateView(GradeBookMixin, ObjectCreateView):
    """
    Класс для создания нового учебного журнала.
    """
    model = Gradebook
    form_class = GradebookForm
    template_name = 'Gradebook/create/grade_book.html'
    properties = ['group_id']

    permission_required = getpermission('Gradebook', 'create_gradebook')

    buttons = [
        Button (
            id = 'to_list',
            name = 'К таблице',
            link_name = getpattern(Gradebook, 'list'),
            permission = getpermission('Gradebook', 'view_gradebook')
        )
    ]

logger = logging.getLogger(__name__)

@permission_required(getpermission('Gradebook', 'download_gradebook'))
def download_report(request, pk):
    try:
        gradebook = get_object_or_404(
            Gradebook.objects.select_related('discipline', 'group__qualification__specialty').prefetch_related('teachers', 'students'),
            pk=pk
        )

        if gradebook.status == Gradebook.STATUS_CHOICE[2][1]:
            gradebook.status = Gradebook.STATUS_CHOICE[3][1]
            gradebook.date_of_closing = timezone.now()
            gradebook.save()

        # Отладка исходных данных
        logger.debug(f"Raw data: name={gradebook.name}, discipline={gradebook.discipline.name}, group={gradebook.group.full_name}, semester={gradebook.semester_number}, pk={pk}")

        # Формирование имени файла
        file_name_base = f"Vedomost_{gradebook.group.full_name}_{gradebook.discipline.name}_{gradebook.semester_number}_{gradebook.type_of_grade_book}"
        # Транслитерация кириллицы в латиницу
        file_name_translit = transliterate(file_name_base)
        # Очистка имени файла от недопустимых символов и замена пробелов на подчеркивания
        file_name_clean = re.sub(r'[^a-zA-Z0-9_]', '_', file_name_translit)[:50]
        file_name = f"{file_name_clean}.docx"

        # Отладка сформированного имени
        logger.debug(f"Generated filename: {file_name}")

        generator = GradebookDocumentGenerator(gradebook)
        buffer = generator.generate_document()

        # Проверка и сброс указателя буфера
        if hasattr(buffer, 'seek') and hasattr(buffer, 'tell'):
            buffer.seek(0)  # Сбрасываем указатель в начало
            buffer_length = len(buffer.getvalue()) if hasattr(buffer, 'getvalue') else 0
            logger.debug(f"Buffer length after seek: {buffer_length}")
        else:
            logger.error("Buffer does not support seek/tell operations")
            raise ValueError("Buffer does not support required operations")

        # Проверка содержимого буфера
        if buffer_length == 0:
            logger.error("Buffer is empty after generation")
            raise ValueError("Generated document buffer is empty")

        response = FileResponse(
            buffer,
            as_attachment=True,
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        # Устанавливаем заголовок для корректной обработки клиентом
        response['Content-Disposition'] = f'attachment; filename="{file_name}"'
        # Разрешаем доступ к заголовку Content-Disposition для JavaScript
        response['Access-Control-Expose-Headers'] = 'Content-Disposition'
        logger.debug(f"Response Content-Disposition: {response['Content-Disposition']}")
        logger.debug(f"Response status: 200")
        return response

    except Exception as e:
        logger.error(f"Error in download_report: {str(e)}", exc_info=True)
        return HttpResponse(f"Произошла ошибка: {str(e)}", status=500)

@permission_required(getpermission('Gradebook', 'view_gradebook'))
def check_and_open_gradebook(request, pk):
    """
    Проверяет заполненность всех обязательных полей ведомости,
    включая наличие связанных преподавателей и студентов.
    """
    gradebook = get_object_or_404(Gradebook, id=pk)

    # Проверка основных полей
    required_fields = [
        gradebook.group_id,  # Проверка ForeignKey через id
        gradebook.name,  # Проверка CharField
        gradebook.discipline_id,  # Проверка ForeignKey дисциплины
        gradebook.semester_number,  # Проверка IntegerField
        gradebook.status,  # Проверка CharField статуса
    ]

    # Проверка текстовых полей на непустые значения
    text_fields_valid = all([
        gradebook.name and gradebook.name.strip(),
        gradebook.status and gradebook.status.strip(),
    ])

    # Проверка связей ManyToMany и ForeignKey
    relations_valid = all([
        gradebook.teachers.exists(),  # Проверка наличия преподавателей
        gradebook.students.exists(),  # Проверка наличия студентов
    ])

    # Комплексная проверка всех условий
    if not all(required_fields) or not text_fields_valid or not relations_valid:
        return redirect('gradebook_detail', pk=pk)

    # Обновление статуса и даты
    gradebook.status = "Открыта"
    gradebook.date_of_opening = timezone.now().date()
    gradebook.save()

    response = redirect('gradebook_detail', pk=pk)
    response['Location'] += '?is_active=true'
    return response


def export_diplom_types(request):
    """Экспортирует таблицу успеваемости и типов диплома в Excel."""
    group_id = request.GET.get('group')
    if not group_id:
        logger.error("No group ID provided for export")
        return HttpResponse("Группа не указана", status=400)

    try:
        group = get_object_or_404(GroupStudents, id=group_id)
        logger.info(f"Exporting diplom types for group {group.full_name}")

        # Получаем дисциплины и общую сумму часов
        disciplines_with_semesters, total_all_hours = get_disciplines_with_last_semester(group)
        if not disciplines_with_semesters:
            logger.warning(f"No disciplines found for group {group.full_name}")
            return HttpResponse("Для выбранной группы нет дисциплин", status=404)

        students = Student.objects.filter(group=group, is_expelled=False, is_in_academ=False)
        semesters_order = sorted({d[1] for d in disciplines_with_semesters})

        # Сортируем дисциплины по семестрам для一致ности
        sorted_disciplines = []
        for semester in semesters_order:
            semester_disciplines = [d for d in disciplines_with_semesters if d[1] == semester]
            sorted_disciplines.extend(sorted(semester_disciplines, key=lambda x: x[0].name))

        # Подготовка оценок с учетом типов ведомостей
        grades_data = {}
        for student in students:
            grades_data[student.id] = {}
            for discipline, semester, _, _ in sorted_disciplines:
                # Получаем все оценки для дисциплины и семестра
                gradebooks = student.gradebookstudents_set.filter(
                    gradebook__discipline=discipline,
                    gradebook__semester_number=semester
                ).order_by('-gradebook__date_of_closing')

                # Логика получения последней актуальной оценки
                grade_obj = gradebooks.first()  # Берём самую последнюю оценку
                grades_data[student.id][(discipline.id, semester)] = grade_obj
                logger.debug(
                    f"Pre-fetched grade for {student.full_name}, {discipline.name}, semester {semester}: "
                    f"{grade_obj.grade if grade_obj else None}, type={grade_obj.gradebook.type_of_grade_book if grade_obj else None}"
                )

        # Создаём Excel-файл
        wb = Workbook()
        ws = wb.active
        ws.title = "Сводная ведомость"

        # Стили
        bold_font = Font(bold=True)
        center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        # Цвета
        FILLS = {
            'Первичная сдача': PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid"),
            'Пересдача': PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid"),
            'Комиссия': PatternFill(start_color="E2D9F3", end_color="9370DB", fill_type="solid"),
            'неуд': PatternFill(start_color="F5C6CB", end_color="F5C6CB", fill_type="solid")
        }
        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
        blue_fill = PatternFill(start_color="2196F3", end_color="2196F3", fill_type="solid")

        # Заголовок
        education_basis = 'После 11 класса' if group.education_base == 'Среднее общее' else 'После 9 класса'
        ws.append([f"СПЕЦИАЛЬНОСТЬ {group.qualification.name} ({education_basis})"])
        ws.append([f"СВОДНАЯ ВЕДОМОСТЬ УСПЕВАЕМОСТИ СТУДЕНТОВ ГРУППЫ {group.full_name}"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2 + len(sorted_disciplines) + 1)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=2 + len(sorted_disciplines) + 1)
        for row in ws[1:2]:
            for cell in row:
                cell.font = bold_font
                cell.alignment = center_align

        # Легенда с цветами
        ws.append([])
        ws.append(["Легенда:"])
        ws.append(["Первичная сдача", "Оценка при первой попытке"])
        ws.append(["Пересдача", "Оценка при пересдаче"])
        ws.append(["Комиссия", "Оценка через комиссию"])
        ws.append(["2 или н/я", "Неудовлетворительная оценка или неявка"])
        ws['A5'].fill = FILLS['Первичная сдача']
        ws['A6'].fill = FILLS['Пересдача']
        ws['A7'].fill = FILLS['Комиссия']
        ws['A8'].fill = FILLS['неуд']
        for row in ws[4:8]:
            for cell in row[:2]:
                cell.border = thin_border
                cell.alignment = center_align
            if row[0].row in [4, 5]:
                row[0].font = bold_font

        # Заголовок таблицы
        header_row = ["№", "ФИО студента"]
        for discipline, disp_semester, hours, display_name in sorted_disciplines:
            header_row.append(f"{discipline.code} {display_name} (Семестр {disp_semester})")
        header_row.append("Тип диплома")
        ws.append([])
        ws.append(header_row)
        header_row_num = ws.max_row

        # Строка с часами
        hours_row = ["", "Часы"]
        for discipline, disp_semester, hours, display_name in sorted_disciplines:
            hours_row.append(hours or 0)
        hours_row.append(total_all_hours or 0)
        ws.append(hours_row)

        # Данные студентов
        total_disciplines = len(sorted_disciplines)
        for idx, student in enumerate(students, 1):
            row = [idx, student.full_name]
            grades = []
            for discipline, semester, hours, display_name in sorted_disciplines:
                grade_obj = grades_data[student.id].get((discipline.id, semester))
                row.append(grade_obj.grade if grade_obj else "-")
                grades.append(grade_obj)

            # Подсчёт значимых оценок
            significant_grades = len([g for g in grades if g and g.grade != "-"])

            # Определение типа диплома
            if significant_grades < total_disciplines * 0.25:
                diploma_type = "Синий"
            else:
                bad_grades = [g for g in grades if g and str(g.grade).lower() in ['2', 'неявка', 'н/я', 'неудовлетворительно']]
                valid_grades = [g for g in grades if g and str(g.grade).lower() not in ['2', 'неявка', 'н/я', 'неудовлетворительно']]
                count_4 = len([g for g in valid_grades if g.grade == '4'])
                total_grades = len(valid_grades)
                diploma_type = "Синий" if bad_grades or (total_grades and count_4 > total_grades * 0.25) else "Красный"

            row.append(diploma_type)
            ws.append(row)

        # Форматирование таблицы
        for row in ws[header_row_num:ws.max_row]:
            for cell in row:
                cell.border = thin_border
                cell.alignment = center_align
        for cell in ws[header_row_num]:
            cell.font = bold_font
        for cell in ws[header_row_num + 1]:
            cell.font = bold_font

        # Установка ширины столбцов
        ws.column_dimensions['A'].width = 5  # №
        ws.column_dimensions['B'].width = 30  # ФИО
        for col in range(3, len(sorted_disciplines) + 3):
            ws.column_dimensions[get_column_letter(col)].width = 10  # Дисциплины
        ws.column_dimensions[get_column_letter(len(sorted_disciplines) + 3)].width = 15  # Тип диплома

        # Применение цветов для оценок и диплома
        for row_idx in range(header_row_num + 2, ws.max_row + 1):
            student = students[row_idx - (header_row_num + 2)]
            for col_idx, (discipline, semester, _, _) in enumerate(sorted_disciplines, start=3):
                cell = ws.cell(row=row_idx, column=col_idx)
                grade = cell.value
                logger.debug(f"Processing cell ({row_idx}, {col_idx}): value={grade}")
                if grade and grade != "-":
                    grade_obj = grades_data[student.id].get((discipline.id, semester))
                    if grade_obj:
                        grade_str = str(grade_obj.grade).lower()
                        type_of_grade = str(grade_obj.gradebook.type_of_grade_book).lower().strip()
                        logger.debug(
                            f"Grade for {student.full_name}, {discipline.name}: grade={grade_str}, type={type_of_grade}")
                        if grade_str in ['2', 'неявка', 'н/я', 'неудовлетворительно']:
                            cell.fill = FILLS['неуд']
                            logger.debug(f"Applied неуд fill for {grade_str}")
                        elif type_of_grade == 'первичная сдача':
                            cell.fill = FILLS['Первичная сдача']
                            logger.debug(f"Applied Первичная сдача fill")
                        elif type_of_grade == 'пересдача':
                            cell.fill = FILLS['Пересдача']
                            logger.debug(f"Applied Пересдача fill")
                        elif type_of_grade == 'комиссия':
                            cell.fill = FILLS['Комиссия']
                            logger.debug(f"Applied Комиссия fill")
                    else:
                        logger.debug(
                            f"No grade_obj found for {student.full_name}, {discipline.name}, semester {semester}")
            # Цвет для типа диплома
            diploma_cell = ws.cell(row=row_idx, column=len(sorted_disciplines) + 3)
            logger.debug(f"Diploma cell ({row_idx}, {len(sorted_disciplines) + 3}): value={diploma_cell.value}")
            if diploma_cell.value == "Красный":
                diploma_cell.fill = red_fill
                diploma_cell.font = Font(color="FFFFFF")
            elif diploma_cell.value == "Синий":
                diploma_cell.fill = blue_fill
                diploma_cell.font = Font(color="FFFFFF")

        # Сохранение файла
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="Diplom_Types_{group.full_name}.xlsx"'
        wb.save(response)
        logger.info(f"Excel file generated for group {group.full_name}")
        return response

    except Exception as e:
        logger.error(f"Error exporting diplom types: {str(e)}", exc_info=True)
        return HttpResponse(f"Ошибка при экспорте: {str(e)}", status=500)